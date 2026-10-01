"""Builds downloadable Excel (.xlsx) and PDF exports of a tracker result.

Like mailer.py, this only ever runs against a result this server fetched
itself (see run_tracker() in app.py) - never client-supplied data.
"""

from io import BytesIO
from xml.sax.saxutils import escape as xml_escape

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

_HEADER_FILL = "1D242A"


def _summary_rows(result):
    rows = [
        ("Line", result.get("line_name", "-")),
        ("Container Number", result.get("container_number", "-")),
    ]
    for field in result.get("summary_fields") or []:
        if isinstance(field, dict) and field.get("label") and field.get("value"):
            rows.append((field["label"], field["value"]))
    return rows


def _event_rows(result):
    columns = result.get("event_columns") or []
    events = result.get("events") or []
    rows = []
    for row in events:
        if isinstance(row, dict):
            rows.append([row.get(col, "") for col in columns])
        else:
            rows.append(list(row))
    return columns, rows


def build_xlsx_bytes(result) -> bytes:
    wb = Workbook()
    header_font = Font(bold=True, color="FFFFFF")

    summary_ws = wb.active
    summary_ws.title = "Summary"
    summary_ws.append(["Field", "Value"])
    for cell in summary_ws[1]:
        cell.font = header_font
        cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
    for label, value in _summary_rows(result):
        summary_ws.append([label, value])
    summary_ws.column_dimensions["A"].width = 24
    summary_ws.column_dimensions["B"].width = 60

    columns, rows = _event_rows(result)
    events_ws = wb.create_sheet("Events")
    if columns:
        events_ws.append(columns)
        for cell in events_ws[1]:
            cell.font = header_font
            cell.fill = PatternFill("solid", fgColor=_HEADER_FILL)
        for row in rows:
            events_ws.append(row)
        for i in range(1, len(columns) + 1):
            events_ws.column_dimensions[get_column_letter(i)].width = 22

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def build_pdf_bytes(result) -> bytes:
    columns, rows = _event_rows(result)

    # Many-column tables (e.g. LDB's 15) don't fit portrait A4: go landscape
    # with tighter margins and a smaller font.
    wide = len(columns) > 8
    buf = BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=landscape(A4) if wide else A4,
        topMargin=1.5 * cm,
        bottomMargin=1.5 * cm,
        leftMargin=1 * cm if wide else 2 * cm,
        rightMargin=1 * cm if wide else 2 * cm,
    )
    styles = getSampleStyleSheet()
    story = []

    title = f"{result.get('line_name', 'Carrier')} — {result.get('container_number', '-')}"
    story.append(Paragraph(xml_escape(title), styles["Title"]))
    story.append(Spacer(1, 12))

    summary_rows = _summary_rows(result)
    if summary_rows:
        story.append(Paragraph("Summary", styles["Heading2"]))
        table_data = [["Field", "Value"]] + [[str(a), str(b)] for a, b in summary_rows]
        t = Table(table_data, colWidths=[150, 350])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_HEADER_FILL}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        story.append(t)
        story.append(Spacer(1, 16))

    if columns and rows:
        story.append(Paragraph("Events", styles["Heading2"]))
        font_size = 6.5 if wide else 8
        cell_style = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=font_size, leading=font_size + 2)
        # Paragraph cells ignore the table's TEXTCOLOR, so the header needs
        # its own white style to be readable on the dark header fill.
        head_style = ParagraphStyle("head", parent=cell_style, textColor=colors.white, fontName="Helvetica-Bold")
        wrapped = [[Paragraph(xml_escape(str(c)), head_style) for c in columns]]
        for row in rows:
            wrapped.append([Paragraph(xml_escape(str(v)) if v else "-", cell_style) for v in row])
        col_width = doc.width / max(len(columns), 1)
        t2 = Table(wrapped, colWidths=[col_width] * len(columns), repeatRows=1)
        t2.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{_HEADER_FILL}")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ]))
        story.append(t2)

    source_url = result.get("source_url")
    if source_url:
        story.append(Spacer(1, 16))
        safe_url = xml_escape(source_url)
        story.append(Paragraph(f'<link href="{safe_url}">View on carrier site</link>', styles["BodyText"]))

    doc.build(story)
    return buf.getvalue()
